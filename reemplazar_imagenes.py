#!/usr/bin/env python3
"""
Reemplaza imágenes en el DOCX con screenshots reales del sistema.
Fase 1: Detecta qué imágenes reemplazar (basado en contexto).
Fase 2: Modifica el ZIP del DOCX en una sola pasada.
"""

import os, re, shutil, zipfile
from lxml import etree
from docx import Document
from docx.oxml.ns import qn

BASE = '/mnt/c/Users/Dell/Documents/cuadroInt/webappcuadrointegral'
SS_DIR = '/tmp/cmi_screenshots'

# ── FIND DOCX ──────────────────────────────────────────────────

for f in os.listdir(BASE):
    if 'CORREGIDO' in f and 'FINAL' not in f:
        DOCX_PATH = os.path.join(BASE, f)
        break

print(f"Documento: {os.path.basename(DOCX_PATH)}")

# ── SCREENSHOT MAPPING ─────────────────────────────────────────
# (screenshot, regex_to_match_in_context)
IMAGE_MAP = [
    ('01_dashboard.png',    'Dashboard'),
    ('02_projects.png',     'Gesti.n de Proyectos'),  # Gesti[óo]n
    ('09_wizard.png',       'Formulario.*Creaci.n'),  # Creaci[óo]n
    ('03_stakeholders.png', 'Registro de stakeholders'),
    ('04_risks.png',        'Gesti.n de riesgos'),
    ('05_resources.png',    'Gesti.n de recursos'),
    ('06_reports.png',      'Reporte de avance'),
    ('07_gantt.png',        'Diagrama de Gantt'),
    ('08_evm_curves.png',   'Gr.ficos de rendimiento'),  # Gr[áa]ficos
]

# ── PHASE 1: Find images to replace ────────────────────────────

print("\n🔍 Fase 1: Buscando imágenes para reemplazar...\n")
doc = Document(DOCX_PATH)

namespaces = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
}

replacements = []  # [(rel_id, old_target, screenshot_path)]

for ss_name, pattern in IMAGE_MAP:
    ss_path = os.path.join(SS_DIR, ss_name)
    if not os.path.exists(ss_path):
        print(f"  ⚠️  Screenshot no encontrado: {ss_name}")
        continue

    found = False
    for i, p in enumerate(doc.paragraphs):
        blips = p._element.findall('.//a:blip', namespaces)
        if not blips:
            continue
        
        # Check context: text of this paragraph OR adjacent paragraphs
        texts = []
        if i > 0:
            texts.append(doc.paragraphs[i-1].text)
        texts.append(p.text)
        if i + 1 < len(doc.paragraphs):
            texts.append(doc.paragraphs[i+1].text)
        context = ' '.join(texts)
        
        if re.search(pattern, context, re.IGNORECASE):
            for blip in blips:
                embed = blip.get(qn('r:embed'))
                if embed:
                    try:
                        rel = doc.part.rels[embed]
                        old_target = rel.target_ref
                        replacements.append((embed, old_target, ss_path, ss_name))
                        print(f"  ✅ {ss_name} → Rel:{embed} ({old_target}) | Contexto: '{p.text.strip()[:50]}'")
                        found = True
                    except:
                        pass
                    break  # One image per screenshot
            if found:
                break
    
    if not found:
        print(f"  ❌ {ss_name} → No se encontró imagen para: '{pattern}'")

print(f"\n   Total a reemplazar: {len(replacements)}")

# ── PHASE 2: Replace in ZIP ────────────────────────────────────

if not replacements:
    print("\n⚠️  No hay imágenes que reemplazar. Saliendo...")
    exit(0)

print("\n🔧 Fase 2: Reemplazando imágenes en el ZIP...")

tmp_path = DOCX_PATH + '.tmp'
shutil.copy2(DOCX_PATH, tmp_path)

# Read all files from zip
with zipfile.ZipFile(tmp_path, 'r') as zin:
    all_files = {}
    for item in zin.infolist():
        all_files[item.filename] = zin.read(item.filename)
    
    # Parse relationships
    rels_data = all_files.get('word/_rels/document.xml.rels', b'')
    if not rels_data:
        print("  ❌ No se encontraron relaciones del documento")
        exit(1)
    
    rels_xml = etree.fromstring(rels_data)
    ns_rel = {'': 'http://schemas.openxmlformats.org/package/2006/relationships'}
    
    # For each replacement, update the relationship target and add the new image
    for embed, old_target, ss_path, ss_name in replacements:
        # Create a new media filename
        new_media_name = f'word/media/ss_{ss_name}'
        
        # Find and update the relationship
        found_rel = False
        for rel_elem in rels_xml:
            rid = rel_elem.get('Id')
            if rid == embed:
                rel_elem.set('Target', new_media_name.replace('word/', ''))
                print(f"    Rel {embed}: {old_target} → {new_media_name}")
                found_rel = True
                break
        
        if not found_rel:
            print(f"    ⚠️  Relación {embed} no encontrada en XML")
            continue
        
        # Read screenshot data
        with open(ss_path, 'rb') as f:
            ss_data = f.read()
        
        # Add the new media file
        all_files[new_media_name] = ss_data
        print(f"    Imagen agregada: {new_media_name} ({len(ss_data)} bytes)")
    
    # Update the rels file in all_files
    new_rels_data = etree.tostring(rels_xml, xml_declaration=True, encoding='UTF-8', standalone=True)
    all_files['word/_rels/document.xml.rels'] = new_rels_data

# Write modified zip
print("\n💾 Escribiendo DOCX modificado...")
with zipfile.ZipFile(tmp_path, 'w', zipfile.ZIP_DEFLATED) as zout:
    for fname, fdata in all_files.items():
        zout.writestr(fname, fdata)

# Replace original
shutil.move(tmp_path, DOCX_PATH)

print(f"\n✅ {len(replacements)} imágenes reemplazadas exitosamente")
print(f"Documento: {DOCX_PATH}")
