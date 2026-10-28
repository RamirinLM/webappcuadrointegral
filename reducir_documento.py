from docx import Document
from docx.oxml.ns import qn
from collections import defaultdict
import re

doc_path = "E4_MERCHÁN_MONTALVÁN_JOSÉ_RAMIRO_FINAL_reducido.docx"
doc = Document(doc_path)

# Find all figure captions
figs_by_num = defaultdict(list)
for idx, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    t_low = t.lower()
    if t_low.startswith('figura'):
        m = re.search(r'figura\s*#?\s*(\d+)', t_low)
        if m:
            num = m.group(1)
            # Determine context
            contexto = 'OTROS'
            for j in range(idx, max(0, idx-100), -1):
                t2 = doc.paragraphs[j].text.lower()
                if 'anexo' in t2 or 'apéndice' in t2 or 'apendice' in t2:
                    contexto = 'ANEXOS'
                    break
                elif 'prueba' in t2:
                    contexto = 'PRUEBAS'
                    break
                elif 'prototipo' in t2:
                    contexto = 'PROTOTIPO'
                    break
                elif ('historia' in t2 and 'usuario' in t2) or ('detalle' in t2 and 'us' in t2):
                    contexto = 'DETALLE_US'
                    break
            figs_by_num[num].append({
                'para_idx': idx,
                'text': t,
                'contexto': contexto
            })

# Decide keeper for each figure
to_replace = []
keepers = {}

for num, items in figs_by_num.items():
    anexos_items = [it for it in items if it['contexto'] == 'ANEXOS']
    if anexos_items:
        keeper = anexos_items[0]
    else:
        keeper = min(items, key=lambda x: x['para_idx'])
    keepers[num] = keeper['para_idx']
    for it in items:
        if it['para_idx'] != keeper['para_idx']:
            to_replace.append((it['para_idx'], num))

# Build set for fast lookup
to_replace_set = set([idx for idx, num in to_replace])
replace_map = {idx: num for idx, num in to_replace}

print(f"Occurrences to replace: {len(to_replace)}")
print(f"Will process paragraphs from end to start to avoid index shifting issues...")

# Process from end to start
replaced = 0
for idx in sorted(list(to_replace_set), reverse=True):
    p = doc.paragraphs[idx]
    fig_num = replace_map[idx]
    # Check if paragraph has drawing
    drawings = p._element.xpath('.//w:drawing')
    if drawings:
        # Remove all drawings
        for d in drawings:
            d.getparent().remove(d)
        # Clear text and add reference
        p.clear()
        run = p.add_run(f'Ver Figura {fig_num}.')
        replaced += 1
    else:
        # No drawing but it's a duplicate caption? Just add reference or modify
        # Try to replace the text
        old_text = p.text
        p.clear()
        p.add_run(f'Ver Figura {fig_num}.')
        replaced += 1

doc.save(doc_path)
print(f"Done. Replaced {replaced} paragraphs.")
